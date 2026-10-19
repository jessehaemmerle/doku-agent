"""Text aus hochgeladenen Dateien gewinnen – lokal, ohne externe Dienste."""

from __future__ import annotations

from pathlib import Path

MAX_BYTES = 20 * 1024 * 1024


class ExtractError(ValueError):
    pass


def extract_text(path: Path) -> str:
    if path.stat().st_size > MAX_BYTES:
        raise ExtractError(f"{path.name}: Datei größer als 20 MB")
    suffix = path.suffix.lower()
    if suffix == ".pdf":
        from pypdf import PdfReader

        reader = PdfReader(str(path))
        return "\n\n".join((page.extract_text() or "") for page in reader.pages).strip()
    if suffix == ".docx":
        import docx

        doc = docx.Document(str(path))
        parts = [p.text for p in doc.paragraphs if p.text.strip()]
        for table in doc.tables:
            for row in table.rows:
                parts.append(" | ".join(cell.text.strip() for cell in row.cells))
        return "\n".join(parts)
    data = path.read_bytes()
    if b"\x00" in data[:4096]:
        raise ExtractError(f"{path.name}: Binärdatei wird nicht unterstützt (nur Text, PDF, DOCX)")
    for encoding in ("utf-8", "cp1252"):
        try:
            return data.decode(encoding)
        except UnicodeDecodeError:
            continue
    return data.decode("utf-8", errors="replace")
